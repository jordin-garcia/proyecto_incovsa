# -*- coding: utf-8 -*-
"""Arma un documento .docx del proyecto a partir de archivos Markdown.

Uso, desde la carpeta entregable-2/:
    python herramientas/armar_documento.py documento-2

La carpeta del documento tiene un documento.json (portada, control de versiones, índices,
orden de los archivos) y los archivos .md de cada sección. El resultado sigue el formato del
Documento 1 y de los lineamientos del curso: carta, márgenes de 2.54 cm, Arial 11, interlineado
1.5, texto justificado, títulos numerados, índices automáticos y número de página.

Los carnés de la portada se leen de entregable-2/equipo.local.json, que no se sube al
repositorio. Formato: [["Nombre completo", "Carné"], ...]

Markdown que se entiende:
    # Título            sección (se numera: 1., 2., ...). Con {-} al final no se numera
    ## Título           subsección (1.1)        ### Título   (1.1.1)
    párrafos con **negrita**, *cursiva*, `código` y [texto](url)
    - viñeta / 1. lista numerada (un nivel de sangría con 4 espacios)
    > nota              párrafo en cursiva con sangría
    ```                 bloque de código
    Tabla: Título {#tab:id}   (línea justo antes de una tabla de Markdown)
    ![Título](ruta.png){#fig:id}
    {{fig:id}} / {{tab:id}}   se reemplaza por «Figura N» / «Tabla N»
    \\salto             salto de página
    \\horizontal / \\vertical   sección nueva en horizontal (tablas anchas) o de regreso a vertical
"""
import copy
import io
import json
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

VERDE = RGBColor(0x14, 0x61, 0x2D)
FONDO_CABECERA = "EAF1EC"
FUENTE = "Arial"


# ------------------------------------------------------------------ estilos
def fuente(estilo_o_run, nombre=FUENTE):
    """Fija la fuente en los cuatro tipos de texto, para que Word no use la del tema."""
    rpr = estilo_o_run.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(k), nombre)
    for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(k)) is not None:
            del rf.attrib[qn(k)]


def preparar_estilos(doc):
    st = doc.styles
    normal = st["Normal"]
    fuente(normal)
    normal.font.size = Pt(11)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_after = Pt(6)
    pf.space_before = Pt(0)

    for nombre, tam, color, alin, antes in (("Heading 1", 16, VERDE, WD_ALIGN_PARAGRAPH.CENTER, 18),
                                            ("Heading 2", 13, RGBColor(0, 0, 0), WD_ALIGN_PARAGRAPH.LEFT, 12),
                                            ("Heading 3", 11.5, RGBColor(0, 0, 0), WD_ALIGN_PARAGRAPH.LEFT, 12)):
        s = st[nombre]
        fuente(s)
        s.font.size = Pt(tam)
        s.font.bold = True
        s.font.italic = False
        s.font.color.rgb = color
        s.paragraph_format.alignment = alin
        s.paragraph_format.space_before = Pt(antes)
        s.paragraph_format.space_after = Pt(8)
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.line_spacing = 1.15

    cap = st["Caption"]
    fuente(cap)
    cap.font.size = Pt(9.5)
    cap.font.bold = False
    cap.font.italic = False
    cap.font.color.rgb = RGBColor(0, 0, 0)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(10)

    for n in ("List Bullet", "List Number", "List Bullet 2", "List Number 2", "Footer"):
        fuente(st[n])
    for n in ("toc 1", "toc 2", "toc 3", "table of figures"):
        try:
            s = st[n]
        except KeyError:
            s = st.add_style(n, 1)
            s.base_style = normal
        fuente(s)
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        s.paragraph_format.line_spacing = 1.15

    cod = st.add_style("Codigo", 1)
    cod.base_style = normal
    fuente(cod, "Consolas")
    cod.font.size = Pt(9)
    cod.paragraph_format.line_spacing = 1.0
    cod.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cod.paragraph_format.space_after = Pt(0)
    cod.paragraph_format.left_indent = Cm(0.5)

    tab = st.add_style("Texto tabla", 1)
    tab.base_style = normal
    tab.font.size = Pt(10)
    tab.paragraph_format.line_spacing = 1.0
    tab.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    tab.paragraph_format.space_after = Pt(0)


# ------------------------------------------------------------------ campos de Word
def campo(parrafo, instruccion, texto="Actualice este campo"):
    """Inserta un campo (TOC, SEQ, PAGE) con un resultado provisional."""
    def r(tag, **attrs):
        run = OxmlElement("w:r")
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        run.append(el)
        return run
    p = parrafo._p
    p.append(r("w:fldChar", **{"w:fldCharType": "begin", "w:dirty": "true"}))
    run = OxmlElement("w:r")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instruccion} "
    run.append(it)
    p.append(run)
    p.append(r("w:fldChar", **{"w:fldCharType": "separate"}))
    res = parrafo.add_run(texto)
    p.append(r("w:fldChar", **{"w:fldCharType": "end"}))
    return res


def sombrear(celda, color):
    tcpr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


ANCHO_UTIL = 16.5      # cm entre márgenes


def anchos_columnas(datos, total=ANCHO_UTIL):
    """Ancho de cada columna según su contenido: la palabra más larga cabe sin partirse y el
    resto del ancho se reparte según la cantidad de texto de cada columna."""
    cols = max(len(f) for f in datos)
    largo, texto = [0] * cols, [0] * cols
    for fila in datos:
        for k, c in enumerate(fila):
            limpio = re.sub(r"[*`]", "", c)
            palabras = limpio.split() or [""]
            largo[k] = max(largo[k], max(len(p) for p in palabras))
            texto[k] += len(limpio)
    minimos = [max(1.3, min(4.6, 0.21 * l + 0.55)) for l in largo]
    sobra = total - sum(minimos)
    if sobra <= 0:
        return [m * total / sum(minimos) for m in minimos]
    peso = [max(t - 4 * l, 0) + 1 for t, l in zip(texto, largo)]
    return [m + sobra * p / sum(peso) for m, p in zip(minimos, peso)]


def fijar_anchos(tabla, anchos_cm):
    tbl = tabla._tbl
    tblpr = tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    grid = tbl.find(qn("w:tblGrid"))
    for gc, w in zip(grid.findall(qn("w:gridCol")), anchos_cm):
        gc.set(qn("w:w"), str(int(w * 567)))
    for fila in tabla.rows:
        for celda, w in zip(fila.cells, anchos_cm):
            celda.width = Cm(w)


def repetir_cabecera(fila):
    trpr = fila._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def salto(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ------------------------------------------------------------------ listas numeradas
class Numeracion:
    """Cada lista numerada empieza en 1 (python-docx, por defecto, continúa la anterior)."""

    def __init__(self, doc):
        self.doc = doc
        self.numbering = doc.part.numbering_part.numbering_definitions._numbering
        estilo = doc.styles["List Number"].element
        self.num_base = int(estilo.pPr.numPr.numId.val)
        abstract = [n for n in self.numbering.findall(qn("w:num")) if int(n.get(qn("w:numId"))) == self.num_base][0]
        self.abstract_id = abstract.find(qn("w:abstractNumId")).get(qn("w:val"))
        self.siguiente = max(int(n.get(qn("w:numId"))) for n in self.numbering.findall(qn("w:num"))) + 1

    def nueva(self):
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(self.siguiente))
        a = OxmlElement("w:abstractNumId")
        a.set(qn("w:val"), self.abstract_id)
        num.append(a)
        lo = OxmlElement("w:lvlOverride")
        lo.set(qn("w:ilvl"), "0")
        so = OxmlElement("w:startOverride")
        so.set(qn("w:val"), "1")
        lo.append(so)
        num.append(lo)
        self.numbering.append(num)
        self.siguiente += 1
        return self.siguiente - 1

    @staticmethod
    def aplicar(parrafo, num_id):
        ppr = parrafo._p.get_or_add_pPr()
        numpr = ppr.find(qn("w:numPr"))
        if numpr is None:
            numpr = OxmlElement("w:numPr")
            ppr.append(numpr)
        for tag, val in (("w:ilvl", "0"), ("w:numId", str(num_id))):
            el = numpr.find(qn(tag))
            if el is None:
                el = OxmlElement(tag)
                numpr.append(el)
            el.set(qn("w:val"), val)


# ------------------------------------------------------------------ texto en línea
TOKENS = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def agregar_texto(parrafo, texto, negrita=False):
    for parte in TOKENS.split(texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**"):
            r = parrafo.add_run(parte[2:-2])
            r.bold = True
        elif parte.startswith("`") and parte.endswith("`"):
            r = parrafo.add_run(parte[1:-1])
            fuente(r, "Consolas")
            r.font.size = Pt(9.5)
            r.bold = negrita or None
        elif parte.startswith("[") and "](" in parte:
            txt, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", parte).groups()
            hipervinculo(parrafo, txt, url)
        elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
            r = parrafo.add_run(parte[1:-1])
            r.italic = True
            r.bold = negrita or None
        else:
            r = parrafo.add_run(parte)
            r.bold = negrita or None


def hipervinculo(parrafo, texto, url):
    if not re.match(r"https?://", url):      # enlaces internos del repositorio: solo el texto
        parrafo.add_run(texto)
        return
    rid = parrafo.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                                 is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, val in (("w:color", "1F5FA8"), ("w:u", "single")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        rpr.append(el)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = texto
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    h.append(run)
    parrafo._p.append(h)


# ------------------------------------------------------------------ Markdown
def leer_bloques(texto):
    """Divide el Markdown en bloques: (tipo, datos)."""
    lineas = texto.replace("\r\n", "\n").split("\n")
    i, bloques = 0, []
    while i < len(lineas):
        ln = lineas[i]
        if not ln.strip():
            i += 1
            continue
        if ln.strip() in ("\\salto", "\\horizontal", "\\vertical"):
            bloques.append((ln.strip()[1:], None))
            i += 1
        elif ln.startswith("```"):
            j = i + 1
            while j < len(lineas) and not lineas[j].startswith("```"):
                j += 1
            bloques.append(("codigo", lineas[i + 1:j]))
            i = j + 1
        elif re.match(r"#{1,3} ", ln):
            nivel = len(ln.split(" ")[0])
            bloques.append(("titulo", (nivel, ln[nivel + 1:].strip())))
            i += 1
        elif ln.startswith("Tabla: "):
            m = re.match(r"Tabla: (.+?)\s*(\{#(tab:[\w-]+)\})?\s*$", ln)
            j = i + 1
            while j < len(lineas) and not lineas[j].strip():      # se admite una línea en blanco antes de la tabla
                j += 1
            if j >= len(lineas) or not lineas[j].lstrip().startswith("|"):
                raise SystemExit(f"«{ln}» debe ir seguida de una tabla")
            filas = []
            while j < len(lineas) and lineas[j].lstrip().startswith("|"):
                filas.append(lineas[j])
                j += 1
            bloques.append(("tabla", (m.group(1), m.group(3), filas)))
            i = j
        elif ln.lstrip().startswith("|"):
            j = i
            filas = []
            while j < len(lineas) and lineas[j].lstrip().startswith("|"):
                filas.append(lineas[j])
                j += 1
            bloques.append(("tabla", (None, None, filas)))
            i = j
        elif ln.startswith("!["):
            m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)\s*(\{#(fig:[\w-]+)\})?", ln)
            bloques.append(("figura", (m.group(1), m.group(2), m.group(4))))
            i += 1
        elif re.match(r"(\s*)([-*]|\d+\.) ", ln):
            items = []
            while i < len(lineas) and (re.match(r"(\s*)([-*]|\d+\.) ", lineas[i]) or
                                       (lineas[i].startswith("  ") and lineas[i].strip() and items)):
                m = re.match(r"(\s*)([-*]|\d+\.) (.*)", lineas[i])
                if m:
                    nivel = 1 if len(m.group(1)) >= 2 else 0
                    items.append([nivel, "num" if m.group(2)[0].isdigit() else "vin", m.group(3)])
                else:
                    items[-1][2] += " " + lineas[i].strip()
                i += 1
            bloques.append(("lista", items))
        elif ln.startswith(">"):
            partes = []
            while i < len(lineas) and lineas[i].startswith(">"):
                partes.append(lineas[i][1:].strip())
                i += 1
            bloques.append(("nota", " ".join(p for p in partes if p)))
        else:
            partes = [ln.strip()]           # la primera línea siempre se toma, para que el ciclo avance
            i += 1
            while i < len(lineas) and lineas[i].strip() and not re.match(r"(#{1,3} |```|\||!\[|>|\s*([-*]|\d+\.) |Tabla: )", lineas[i]):
                partes.append(lineas[i].strip())
                i += 1
            bloques.append(("parrafo", " ".join(partes)))
    return bloques


def celdas(fila):
    return [c.strip() for c in fila.strip().strip("|").split("|")]


class Armador:
    def __init__(self, carpeta, config):
        self.carpeta = carpeta
        self.config = config
        self.doc = Document()
        preparar_estilos(self.doc)
        self.numeracion = Numeracion(self.doc)
        self.contador = [0, 0, 0]
        self.refs = {}
        self.num_fig = 0
        self.num_tab = 0

    # -------------------------------------------------------------- página
    def pagina(self):
        s = self.doc.sections[0]
        s.page_width, s.page_height = Cm(21.59), Cm(27.94)
        for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(s, lado, Cm(2.54))
        s.different_first_page_header_footer = bool(self.config.get("portada", True))
        p = s.footer.paragraphs[0]
        p.style = self.doc.styles["Footer"]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        campo(p, "PAGE", "1")
        ajustes = self.doc.settings.element
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        ajustes.append(uf)

    # -------------------------------------------------------------- portada
    def centrado(self, texto, tam=12, negrita=True, antes=0, despues=4):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(antes)
        p.paragraph_format.space_after = Pt(despues)
        r = p.add_run(texto)
        r.bold = negrita
        r.font.size = Pt(tam)
        return p

    def portada(self):
        c = self.config
        self.centrado("UNIVERSIDAD RAFAEL LANDÍVAR", 18, antes=60)
        for t in ("CAMPUS SAN ALBERTO HURTADO, S.J. DE QUETZALTENANGO", "FACULTAD DE INGENIERÍA",
                  "INGENIERÍA EN INFORMÁTICA Y SISTEMAS", "PROGRAMACIÓN WEB", "ING. MANUEL ROJAS"):
            self.centrado(t)
        self.centrado(c["titulo_documento"], 14, antes=40, despues=8)
        self.centrado("Sitio web institucional autoadministrable y sistema de control logístico de bodega INCOVSA",
                      16, despues=8)
        self.centrado("Empresa Beneficiaria: Industria del Cuero Covadonga S. A.", 12, despues=24)
        self.centrado("Integrantes del equipo", 12, antes=12)
        equipo = self.equipo()
        t = self.doc.add_table(rows=1, cols=2)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for celda, txt in zip(t.rows[0].cells, ("Nombre", "Carné")):
            self.texto_celda(celda, txt, True)
            sombrear(celda, FONDO_CABECERA)
        for nombre, carne in equipo:
            fila = t.add_row().cells
            self.texto_celda(fila[0], nombre)
            self.texto_celda(fila[1], carne)
        fijar_anchos(t, [8.5, 3.0])
        self.centrado(c["lugar_fecha"], 12, antes=36)
        salto(self.doc)

    def equipo(self):
        ruta = os.path.join(self.carpeta, "..", "equipo.local.json")
        if not os.path.exists(ruta):
            raise SystemExit("Falta entregable-2/equipo.local.json con los nombres y carnés del equipo "
                             '([["Nombre", "Carné"], ...]). No se sube al repositorio.')
        return json.load(io.open(ruta, encoding="utf-8"))

    def control_versiones(self):
        self.centrado("TABLA DE CONTROL DE VERSIONES", 12, despues=10)
        t = self.doc.add_table(rows=1, cols=4)
        t.style = "Table Grid"
        for celda, txt in zip(t.rows[0].cells, ("Versión", "Fecha", "Responsable", "Descripción del cambio")):
            self.texto_celda(celda, txt, True)
            sombrear(celda, FONDO_CABECERA)
        for v in self.config["versiones"]:
            fila = t.add_row().cells
            for celda, txt in zip(fila, v):
                self.texto_celda(celda, txt)
        fijar_anchos(t, [1.7, 2.3, 4.2, 8.3])
        salto(self.doc)

    def indices(self):
        for titulo, instr in (("ÍNDICE GENERAL", 'TOC \\o "1-3" \\h \\z \\u'),
                              ("ÍNDICE DE FIGURAS", 'TOC \\h \\z \\c "Figura"'),
                              ("ÍNDICE DE TABLAS", 'TOC \\h \\z \\c "Tabla"')):
            self.centrado(titulo, 12, despues=10)
            p = self.doc.add_paragraph()
            campo(p, instr, "Abra el documento en Word y acepte actualizar los campos para ver este índice.")
            salto(self.doc)

    # -------------------------------------------------------------- contenido
    def texto_celda(self, celda, texto, negrita=False):
        p = celda.paragraphs[0]
        p.style = self.doc.styles["Texto tabla"]
        agregar_texto(p, self.referencias(texto), negrita)

    def referencias(self, texto):
        return re.sub(r"\{\{((?:fig|tab):[\w-]+)\}\}", lambda m: self.refs.get(m.group(1), "[referencia sin definir]"), texto)

    def numerar_referencias(self, bloques):
        """Primera pasada: asigna «Figura N» y «Tabla N» a cada id, para poder citarlos antes."""
        f = t = 0
        for tipo, datos in bloques:
            if tipo == "figura":
                f += 1
                if datos[2]:
                    self.refs[datos[2]] = f"Figura {f}"
            elif tipo == "tabla" and datos[0]:
                t += 1
                if datos[1]:
                    self.refs[datos[1]] = f"Tabla {t}"

    def leyenda(self, tipo, titulo):
        p = self.doc.add_paragraph(style="Caption")
        p.add_run(f"{tipo} ")
        n = self.num_fig if tipo == "Figura" else self.num_tab
        campo(p, f"SEQ {tipo} \\* ARABIC", str(n))
        agregar_texto(p, f". {titulo}")
        return p

    def titulo(self, nivel, texto):
        sin_numero = texto.endswith("{-}")
        texto = texto.replace("{-}", "").strip()
        if not sin_numero and self.config.get("numerar", True):
            self.contador[nivel - 1] += 1
            for k in range(nivel, 3):
                self.contador[k] = 0
            texto = ".".join(str(x) for x in self.contador[:nivel]) + (". " if nivel == 1 else " ") + texto
        if nivel == 1:
            texto = texto.upper()
            # Página nueva solo para anexos y referencias (o en cada sección, si el documento lo pide)
            if self.ya_hay_contenido and not self.recien_seccion and (sin_numero or self.config.get("pagina_por_seccion", False)):
                salto(self.doc)
        self.doc.add_heading(texto, level=nivel)
        self.ya_hay_contenido = True
        self.recien_seccion = False

    def orientacion(self, horizontal):
        """Sección nueva en horizontal (para tablas anchas) o de regreso a vertical."""
        s = self.doc.add_section(WD_SECTION.NEW_PAGE)
        s.different_first_page_header_footer = False
        ancho, alto = Cm(21.59), Cm(27.94)
        s.orientation = 1 if horizontal else 0
        s.page_width, s.page_height = (alto, ancho) if horizontal else (ancho, alto)
        self.ancho_util = (27.94 if horizontal else 21.59) - 2 * 2.54
        self.recien_seccion = True

    def tabla(self, titulo, filas):
        filas = [f for f in filas if not re.match(r"^\s*\|?\s*:?-{3,}", f)]
        datos = [celdas(f) for f in filas]
        if titulo:
            self.num_tab += 1
            self.leyenda("Tabla", titulo).paragraph_format.keep_with_next = True
        cols = max(len(r) for r in datos)
        t = self.doc.add_table(rows=0, cols=cols)
        t.style = "Table Grid"
        for k, fila in enumerate(datos):
            nueva = t.add_row()
            no_partir = OxmlElement("w:cantSplit")          # la fila no se parte entre dos páginas
            nueva._tr.get_or_add_trPr().append(no_partir)
            cells = nueva.cells
            for celda, txt in zip(cells, fila + [""] * (cols - len(fila))):
                self.texto_celda(celda, txt, k == 0)
                if k == 0:
                    sombrear(celda, FONDO_CABECERA)
        repetir_cabecera(t.rows[0])
        fijar_anchos(t, anchos_columnas(datos, self.ancho_util))
        self.recien_seccion = False
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def figura(self, titulo, ruta):
        ruta_abs = os.path.normpath(os.path.join(self.carpeta, ruta))
        if not os.path.exists(ruta_abs):
            raise SystemExit(f"No existe la figura {ruta_abs}")
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.line_spacing = 1.0
        p.add_run().add_picture(ruta_abs, width=Cm(16))
        self.num_fig += 1
        self.leyenda("Figura", titulo)

    def lista(self, items):
        num_id = None
        for nivel, tipo, texto in items:
            if tipo == "num":
                estilo = "List Number" if nivel == 0 else "List Number 2"
                p = self.doc.add_paragraph(style=estilo)
                if nivel == 0:
                    if num_id is None:
                        num_id = self.numeracion.nueva()
                    Numeracion.aplicar(p, num_id)
            else:
                p = self.doc.add_paragraph(style="List Bullet" if nivel == 0 else "List Bullet 2")
            p.paragraph_format.space_after = Pt(2)
            agregar_texto(p, self.referencias(texto))

    def contenido(self):
        self.ya_hay_contenido = False
        self.recien_seccion = False
        self.ancho_util = ANCHO_UTIL
        todos = []
        for nombre in self.config["secciones"]:
            todos += leer_bloques(io.open(os.path.join(self.carpeta, nombre), encoding="utf-8").read())
        self.numerar_referencias(todos)
        for tipo, datos in todos:
            if tipo == "titulo":
                self.titulo(*datos)
            elif tipo == "parrafo":
                agregar_texto(self.doc.add_paragraph(), self.referencias(datos))
            elif tipo == "lista":
                self.lista(datos)
            elif tipo == "nota":
                p = self.doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.8)
                agregar_texto(p, self.referencias(datos))
                for r in p.runs:
                    r.italic = True
            elif tipo == "codigo":
                for ln in datos or [""]:
                    self.doc.add_paragraph(ln if ln else " ", style="Codigo")
                self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
            elif tipo == "tabla":
                self.tabla(datos[0], datos[2])
            elif tipo == "figura":
                self.figura(datos[0], datos[1])
            elif tipo == "salto":
                salto(self.doc)
            elif tipo in ("horizontal", "vertical"):
                self.orientacion(tipo == "horizontal")

    def armar(self):
        self.pagina()
        if self.config.get("portada", True):
            self.portada()
        if self.config.get("control_versiones", True):
            self.control_versiones()
        if self.config.get("indices", True):
            self.indices()
        self.contenido()
        salida = os.path.join(self.carpeta, "..", self.config["archivo_salida"])
        self.doc.save(salida)
        return os.path.normpath(salida), self.num_fig, self.num_tab


def main():
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    carpeta = os.path.abspath(sys.argv[1])
    config = json.load(io.open(os.path.join(carpeta, "documento.json"), encoding="utf-8"))
    salida, nf, nt = Armador(carpeta, config).armar()
    print(f"Listo: {salida} · {nf} figuras · {nt} tablas")


if __name__ == "__main__":
    main()
