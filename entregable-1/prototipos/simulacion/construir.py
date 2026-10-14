# -*- coding: utf-8 -*-
"""Arma la simulación navegable de INCOVSA en un solo archivo HTML.

Lee el código de src/, los recursos visuales y los datos de prueba que entregó
la empresa (docs_incovsa/, fuera del repositorio) y escribe
../prototipo-simulacion-incovsa.html, que se abre con doble clic.

El archivo resultante contiene datos reales de la empresa (productos y Peso de
Conversión, nombres de clientes, imágenes y documentos técnicos). No se debe
subir al repositorio (RNF-10).

Requiere Python 3 con Pillow, numpy y openpyxl.
Uso, desde esta carpeta:  python construir.py
"""
import base64, io, json, os, re, collections, unicodedata
from PIL import Image
import numpy as np
import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.abspath(os.path.join(AQUI, "..", "..", ".."))
EMPRESA = os.path.join(RAIZ, "docs_incovsa")
SRC = os.path.join(AQUI, "src")
SALIDA = os.path.join(AQUI, "..", "prototipo-simulacion-incovsa.html")


# ------------------------------------------------------------------ datos
# Reglas de specs/13-datos-maestros.md (decisiones del 11/10/2026)
EXCLUIDOS = {"METABISULFITO DE SODIO II", "SAL (INVENTARIO II)", "TRICLORO GRANULAR-TON/0050KG-9127"}
CORREGIDOS = {
    "SULFATO DE ALUMINIO TUR-SAC/0025KG-5449": "SULFATO DE ALUMINIO-TUR-SAC/0025KG-5449",
    "SULFATO DE ALUMINIO TUR-SAC/0025KG-9232": "SULFATO DE ALUMINIO-TUR-SAC/0025KG-9232",
    "SULFURO DE SODIO TUR-SAC/0025KG-9131": "SULFURO DE SODIO-TUR-SAC/0025KG-9131",
    "SULFATO DE SODIO CHI-SAC/0050KG-4765": "SULFATO DE SODIO-CHI-SAC/0050KG-4765",
    "SULFATO DE SODIO CHI-SAC/0050KG-0823": "SULFATO DE SODIO-CHI-SAC/0050KG-0823",
    "SODA CAUSTICA-CHISAC/0025KG-1240": "SODA CAUSTICA-CHI-SAC/0025KG-1240",
    # P-88 (12/10/2026): se unen a INCO BLACK VSF 600% (IND-CAJ/0025KG) y a SODA ASH (SOL-SAC/0025KG)
    "INCOBLACK VSF 600%-IND-CAJ/0025KG-4034": "INCO BLACK VSF 600%-IND-CAJ/0025KG-4034",
    "SODa ASH-SOL-SAC/0025KG-5813": "SODA ASH-SOL-SAC/0025KG-5813",
    "SODA ASH-SOL-SAC/SAC/0025KG-2479": "SODA ASH-SOL-SAC/0025KG-2479",
}
PESO_FIJO = {"METANOL (GUA-TON/0055GL)": 164}


def nombre_odoo(codigo):
    """NOMBRE-ORIGEN-ENVASE/PRESENTACIÓN[-CALIDAD]-LOTE → NOMBRE (ORIGEN-ENVASE/PRESENTACIÓN[-CALIDAD]).
    El origen es el segmento de 3 letras mayúsculas que va justo antes del envase/presentación
    (el segmento con «/»); así no se confunden nombres como INCO-EAL o P/600. Los productos por
    unidad no traen «/»: ahí el origen es el primer segmento de 3 letras después del nombre."""
    s = codigo.split("-")[:-1]
    tres = lambda k: re.fullmatch(r"[A-Z]{3}", s[k].strip())
    k = next((k for k in range(1, len(s) - 1) if tres(k) and "/" in s[k + 1]), None)
    if k is None:
        k = next((k for k in range(1, len(s)) if tres(k)), None)
    if k is None:
        return None
    return "-".join(s[:k]).strip() + " (" + "-".join(x.strip() for x in s[k:]) + ")"


def datos():
    """Departamentos y municipios de Guatemala (INE) y, de Picking II, productos (nombre en formato de
    Odoo, Peso de Conversión y lotes) y clientes con su sector. Desde el 13/10/2026 el sistema no tiene
    catálogos de productos ni de clientes (specs/13): la simulación los usa solo para armar las rutas y
    los pedidos de ejemplo con que arranca. Lo que se carga en Facturación viene en el archivo de Odoo."""
    picking = openpyxl.load_workbook(os.path.join(EMPRESA, "PROYECTO PICKING", "Picking II.xlsm"),
                                     read_only=True, data_only=True)
    pesos, lotes, vistos = collections.OrderedDict(), {}, set()
    for i, r in enumerate(picking["Base Datos Productos"].iter_rows(values_only=True)):
        if i == 0 or not r[0]:
            continue
        codigo = str(r[0]).strip()
        if codigo in vistos:          # ACIDO ACETICO-GUA-GAR/0030KG-GLACI-8566 viene duplicado
            continue
        vistos.add(codigo)
        if codigo in EXCLUIDOS:
            continue
        nombre = nombre_odoo(CORREGIDOS.get(codigo, codigo))
        if not nombre:
            raise SystemExit(f"Código sin formato reconocible: {codigo}")
        pesos.setdefault(nombre, set()).add(r[2])
        lotes.setdefault(nombre, []).append(codigo.split("-")[-1].strip())   # el sufijo del código es el lote de Odoo
    productos = []
    for nombre, ps in pesos.items():
        if nombre in PESO_FIJO:
            ps = {PESO_FIJO[nombre]}
        if len(ps) != 1:
            raise SystemExit(f"{nombre} tiene dos Pesos de Conversión distintos: {ps}")
        productos.append([nombre, next(iter(ps)), lotes[nombre][:3]])
    if len(productos) != 336:
        raise SystemExit(f"Se esperaban 336 productos y salieron {len(productos)}. Revise las reglas de specs/13.")
    # Clave normalizada única (specs/16, decisión 13): mayúsculas, sin acentos, sin puntos ni espacios
    clave = lambda n: re.sub(r"[\s.]", "", "".join(c for c in unicodedata.normalize("NFD", n.upper())
                                                   if unicodedata.category(c) != "Mn"))
    repetidas = [k for k, v in collections.Counter(clave(p[0]) for p in productos).items() if v > 1]
    if repetidas:
        raise SystemExit(f"Productos con la misma clave normalizada: {repetidas}")

    clave_sector = lambda s: re.sub(r"[\s,.]", "", s).lower()
    sectores, clientes = [], []
    for i, r in enumerate(picking["Base Datos Clientes"].iter_rows(values_only=True)):
        if i == 0 or not r[1]:
            continue
        nombre = str(r[1]).strip()
        sector = str(r[2]).strip() if r[2] else ""
        if sector not in sectores:
            sectores.append(sector)
        # Duplicado exacto (mismo nombre y mismo sector): se deja el primero del archivo
        if any(c[0] == nombre and clave_sector(sectores[c[1]]) == clave_sector(sector) for c in clientes):
            continue
        clientes.append([nombre, sectores.index(sector)])
    clientes = [[n, sectores[i]] for n, i in clientes]
    if len(clientes) != 314:
        raise SystemExit(f"Se esperaban 314 clientes y salieron {len(clientes)}. Revise las reglas de specs/13.")

    guatemala = json.load(io.open(os.path.join(AQUI, "datos", "guatemala.json"), encoding="utf-8"))
    deps = [[d["n"], [m["n"] for m in d["municipios"]]] for d in guatemala["departamentos"]]
    nm = sum(len(d[1]) for d in deps)
    if len(deps) != 22 or nm != 340:
        raise SystemExit(f"Se esperaban 22 departamentos y 340 municipios; hay {len(deps)} y {nm}.")

    print(f"  productos {len(productos)} · clientes {len(clientes)} · sectores {len(sectores)}"
          f" · departamentos {len(deps)} · municipios {nm}")
    return {"productos": productos, "clientes": clientes, "guatemala": deps}


# ------------------------------------------------------------------ imágenes
def webp(img, calidad=76, **kw):
    b = io.BytesIO()
    img.save(b, "WEBP", quality=calidad, method=6, **kw)
    return "data:image/webp;base64," + base64.b64encode(b.getvalue()).decode()


def logo(ruta, caja):
    """Recorta el logotipo del JPG y vuelve transparente el fondo blanco.
    Devuelve la versión a color y la versión blanca para fondos oscuros."""
    im = np.array(Image.open(ruta).convert("RGB").crop(caja)).astype(float)
    alfa = np.clip((255 - im.min(axis=2)) * 1.5, 0, 255)
    alfa[alfa < 18] = 0
    a = np.maximum(alfa / 255, 1e-6)[..., None]
    rgb = np.clip((im - (1 - a) * 255) / a, 0, 255)
    color = Image.fromarray(np.dstack([rgb, alfa]).astype(np.uint8), "RGBA")
    color = color.crop(color.getbbox())
    blanco = np.array(color)
    blanco[..., :3] = 255
    return webp(color, 92), webp(Image.fromarray(blanco, "RGBA"), 92)


def recursos():
    info = os.path.join(EMPRESA, "PAGINA WEB", "INFORMACION EMPRESA")
    a = {}
    a["logo"], a["logoBlanco"] = logo(os.path.join(info, "INCOVSA QUIMICOS.jpg"), (262, 195, 760, 392))
    # Solo la parte de la fotografía, sin el texto de la imagen
    a["fotoEquipo"] = webp(Image.open(os.path.join(info, "VISION.jpg")).convert("RGB").crop((0, 0, 1920, 445)), 72)
    a["fotoReunion"] = webp(Image.open(os.path.join(info, "MISION.jpg")).convert("RGB").crop((0, 0, 1440, 430)), 72)
    # Logotipo de Inquiserca, recibido el 13/10/2026 (specs/05): a color, con el fondo blanco transparente
    a["inquiserca"], _ = logo(os.path.join(info, "Inquiserca.jpeg"), (0, 0, 1280, 350))

    catalogo = os.path.join(EMPRESA, "PAGINA WEB", "CATALOGO")
    a["catalogo"] = []
    for f in sorted(os.listdir(catalogo)):
        im = Image.open(os.path.join(catalogo, f)).convert("RGB").resize((560, 560), Image.LANCZOS)
        a["catalogo"].append(webp(im, 74))

    docs = os.path.join(EMPRESA, "PAGINA WEB", "DOCUMENTOS TECNICOS")
    a["pdf"] = {}
    for clave, f in [("ft", "FT SODA CAUSTICA MX.pdf"), ("hs", "HOJA DE SEGURIDAD SODA CAUSTICA MX.pdf"),
                     ("coa", "CAUSTICA MX CAB250526 SKU 4139.pdf")]:
        a["pdf"][clave] = base64.b64encode(open(os.path.join(docs, f), "rb").read()).decode()
    print(f"  {len(a['catalogo'])} imágenes de catálogo · logotipos de INCOVSA e Inquiserca · 2 fotografías · 3 PDF")
    return a


# ------------------------------------------------------------------ armado
def leer(nombre):
    return io.open(os.path.join(SRC, nombre), encoding="utf-8").read()


def a_script(obj):
    return json.dumps(obj, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")


print("Armando la simulación…")
html = leer("plantilla.html")
html = html.replace("/*ESTILOS*/", leer("estilos.css"))
js = "\n".join(leer(f) for f in ["base.js", "sitio.js", "panel.js", "picking.js", "tablero.js", "inicio.js"])
html = html.replace("/*RECURSOS*/", "window.RECURSOS=" + a_script(recursos()) + ";")
html = html.replace("/*DATOS*/", "window.DATOS=" + a_script(datos()) + ";")
html = html.replace("/*APLICACION*/", js.replace("</script", "<\\/script"))
io.open(SALIDA, "w", encoding="utf-8").write(html)
print(f"Listo: {os.path.abspath(SALIDA)} · {os.path.getsize(SALIDA) // 1024} KB")
